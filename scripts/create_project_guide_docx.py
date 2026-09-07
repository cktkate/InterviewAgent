"""Convert the InterviewAgent Markdown guide into a polished DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "CBD5E1"
CODE_FILL = "F3F4F6"
WHITE = "FFFFFF"
PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    latin = "Consolas" if mono else "Calibri"
    east_asia = "Microsoft YaHei"
    run.font.name = latin
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def choose_widths(rows):
    cols = max(len(row) for row in rows)
    if cols == 2:
        first = max(len(row[0]) if row else 0 for row in rows)
        second = max(len(row[1]) if len(row) > 1 else 0 for row in rows)
        ratio = min(0.40, max(0.20, first / max(1, first + second)))
        left = int(PAGE_WIDTH_DXA * ratio)
        return [left, PAGE_WIDTH_DXA - left]
    weights = []
    for idx in range(cols):
        length = max(len(row[idx]) if idx < len(row) else 0 for row in rows)
        weights.append(max(8, min(36, length)))
    total = sum(weights)
    widths = [int(PAGE_WIDTH_DXA * w / total) for w in weights]
    widths[-1] += PAGE_WIDTH_DXA - sum(widths)
    return widths


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("1F2937")
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    pf.widow_control = True

    configs = {
        "Title": (28, INK, 0, 8),
        "Subtitle": (12, MUTED, 0, 8),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in configs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


INLINE_PATTERN = re.compile(r"(\*\*.*?\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text, *, base_size=11, base_color="1F2937"):
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=INK)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(9, base_size - 0.5), mono=True, color=DARK_BLUE)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), CODE_FILL)
            run._element.get_or_add_rPr().append(shd)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, italic=True, color=base_color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=base_size, color=base_color)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(84)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AI AGENT · PROJECT GUIDE")
    set_run_font(run, size=11, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("InterviewAgent\n")
    subtitle_run = title.add_run("项目总结与面试复习手册")
    set_run_font(subtitle_run, size=20, bold=True, color=INK)
    for run in title.runs[:-1]:
        set_run_font(run, size=30, bold=True, color=INK)

    sub = doc.add_paragraph(style="Subtitle")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(sub, "面向 AI Agent 开发 / AI 应用工程师 / RAG 应用工程师", base_size=12, base_color=MUTED)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_before = Pt(16)
    line.paragraph_format.space_after = Pt(28)
    run = line.add_run("LANGGRAPH  ·  HYBRID RAG  ·  MEMORY  ·  TOOL RUNTIME  ·  EVALUATION")
    set_run_font(run, size=9.5, bold=True, color=BLUE)

    p = add_box_paragraph(doc, LIGHT_GRAY, BORDER)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "以当前代码实现为准：区分已实现能力、实验指标与生产化改进。", base_size=10.5, base_color=INK)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(88)
    add_inline(meta, "Interview preparation edition · 2026", base_size=10, base_color=MUTED)
    doc.add_page_break()


def add_running_furniture(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("InterviewAgent · AI Agent 项目复习手册")
    set_run_font(hr, size=9, bold=True, color=MUTED)
    hp_pr = hp._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), BORDER)
    borders.append(bottom)
    hp_pr.append(borders)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_box_paragraph(doc, fill, border_color):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), border_color)
        borders.append(node)
    p_pr.append(borders)
    return p


def add_callout(doc, text):
    p = add_box_paragraph(doc, LIGHT_GRAY, BORDER)
    add_inline(p, text, base_size=10.5, base_color=INK)


def add_code_block(doc, lines):
    p = add_box_paragraph(doc, CODE_FILL, BORDER)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, size=9, mono=True, color="263238")
        if idx != len(lines) - 1:
            run.add_break()


def add_table(doc, rows):
    widths = choose_widths(rows)
    table = doc.add_table(rows=len(rows), cols=len(widths))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for r_idx, source_row in enumerate(rows):
        row = table.rows[r_idx]
        for c_idx, value in enumerate(source_row):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            add_inline(
                p,
                value,
                base_size=9.5 if len(rows) > 5 else 10,
                base_color=INK,
            )
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
        if r_idx == 0:
            set_repeat_table_header(row)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def strip_table_cell(value):
    value = value.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]
    return [item.strip() for item in re.split(r"(?<!\\)\|", value)]


def convert_markdown(markdown, output):
    doc = Document()
    configure_styles(doc)
    add_running_furniture(doc)
    add_cover(doc)

    lines = markdown.read_text(encoding="utf-8").splitlines()
    # The DOCX has its own cover; skip the Markdown title and opening note.
    idx = 0
    if lines and lines[0].startswith("# "):
        idx = 1
    while idx < len(lines) and (not lines[idx].strip() or lines[idx].startswith(">")):
        idx += 1

    in_code = False
    code_lines = []
    while idx < len(lines):
        raw = lines[idx]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue
        if not stripped:
            idx += 1
            continue

        if stripped.startswith("|") and idx + 1 < len(lines):
            next_line = lines[idx + 1].strip()
            if next_line.startswith("|") and re.match(r"^\|?[\s:|-]+\|?$", next_line):
                rows = [strip_table_cell(stripped)]
                idx += 2
                while idx < len(lines) and lines[idx].strip().startswith("|"):
                    rows.append(strip_table_cell(lines[idx].strip()))
                    idx += 1
                add_table(doc, rows)
                continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = max(1, len(heading.group(1)) - 1)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, heading.group(2), base_size={1: 16, 2: 13, 3: 12}[level], base_color=BLUE if level < 3 else DARK_BLUE)
            idx += 1
            continue

        if stripped.startswith(">"):
            add_callout(doc, stripped.lstrip("> "))
            idx += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            idx += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
            idx += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        idx += 1

    if in_code and code_lines:
        add_code_block(doc, code_lines)

    props = doc.core_properties
    props.title = "InterviewAgent 项目总结与面试复习手册"
    props.subject = "AI Agent 开发与AI应用工程师面试复习"
    props.author = "InterviewAgent"
    props.keywords = "LangGraph, AI Agent, RAG, Redis, MySQL, WebSocket, MCP"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    convert_markdown(args.input, args.output)


if __name__ == "__main__":
    main()
